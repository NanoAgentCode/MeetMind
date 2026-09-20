from io import BytesIO

from docx import Document
from docx.shared import Pt

from .models import Meeting


def markdown(meeting: Meeting) -> str:
    assert meeting.minutes
    m = meeting.minutes
    section = lambda title, items: f"## {title}\n\n" + "\n".join(f"- {item}" for item in items)
    return "\n\n".join([
        f"# {m.title}",
        f"> 原始文件：{meeting.filename}",
        f"## 会议摘要\n\n{m.summary}",
        section("关键要点", m.key_points),
        section("会议决策", m.decisions),
        section("行动事项", m.action_items),
    ]) + "\n"


def docx_bytes(meeting: Meeting) -> bytes:
    assert meeting.minutes
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    doc.add_heading(meeting.minutes.title, level=0)
    doc.add_paragraph(f"原始文件：{meeting.filename}")
    doc.add_heading("会议摘要", level=1)
    doc.add_paragraph(meeting.minutes.summary)
    for heading, items in (
        ("关键要点", meeting.minutes.key_points),
        ("会议决策", meeting.minutes.decisions),
        ("行动事项", meeting.minutes.action_items),
    ):
        doc.add_heading(heading, level=1)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()

